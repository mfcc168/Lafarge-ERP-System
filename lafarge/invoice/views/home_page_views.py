import calendar
import logging
import re
from collections import defaultdict
from datetime import datetime

from django.contrib.admin.views.decorators import staff_member_required
from django.db.models import Q, Sum
from django.db.models.functions import ExtractMonth, ExtractYear
from django.http import HttpResponse, HttpResponseBadRequest, JsonResponse
from django.shortcuts import render
from django.utils.timezone import localdate
from django.utils.timezone import make_aware
from django.utils.timezone import now
from django.views.decorators.http import require_GET
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..models import Invoice

logger = logging.getLogger(__name__)


def _delivery_invoices(delivery_date):
    """Include deliveries on the given date and invoices awaiting delivery."""
    return Invoice.objects.filter(
        Q(delivery_date=delivery_date) | Q(delivery_date__isnull=True),
    ).select_related(
        'customer', 'salesman', 'deliveryman',
    ).prefetch_related('invoiceitem_set__product').order_by('pk')


def _delivery_items(invoice, total_quantities=False):
    """Group product lots, optionally totaling quantities for the Excel export."""
    grouped_items = defaultdict(list)
    for item in invoice.invoiceitem_set.all():
        if item.product:
            clean_name = re.sub(r"\s*\(Lot\s*no\.?:?\s*[A-Za-z0-9-]+\)", "", item.product.name)
            grouped_items[clean_name].append(item.quantity)
    if total_quantities:
        return [f"{name} {sum(quantities):.1f}" for name, quantities in grouped_items.items()]
    return [f"{name} ({' + '.join(str(quantity) for quantity in quantities)})"
            for name, quantities in grouped_items.items()]


@staff_member_required
@require_GET
def export_today_deliveries(request):
    """Export selected deliveries with one row per product and summed quantities."""
    raw_ids = request.GET.getlist('invoice_ids')
    if not raw_ids:
        return HttpResponseBadRequest('Select at least one delivery to export.')
    try:
        invoice_ids = {int(value) for value in raw_ids}
    except (TypeError, ValueError):
        return HttpResponseBadRequest('Invalid delivery selection.')
    if any(value <= 0 or value > 9223372036854775807 for value in invoice_ids):
        return HttpResponseBadRequest('Invalid delivery selection.')

    today = localdate()
    invoices = list(_delivery_invoices(today).filter(pk__in=invoice_ids))
    if not invoices:
        return HttpResponseBadRequest('No selected invoices are delivered today or awaiting delivery.')

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Today's deliveries"
    sheet.append(['Customer', 'Address', 'Items'])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type='solid', fgColor='F2F4F7')

    for invoice in invoices:
        # Keep a selected delivery even if it does not have any items yet.
        for item in _delivery_items(invoice, total_quantities=True) or ['']:
            sheet.append([str(invoice.customer), invoice.customer.address or '', item])
            for cell in sheet[sheet.max_row]:
                # Customer and product text must never become Excel formulas.
                cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)

    sheet.column_dimensions['A'].width = 36
    sheet.column_dimensions['B'].width = 60
    sheet.column_dimensions['C'].width = 60
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="todays-deliveries-{today.isoformat()}.xlsx"'
    workbook.save(response)
    return response


@staff_member_required
def home(request):
    """Display today's deliveries, undelivered invoices, and pending deposits."""
    start_date = make_aware(datetime(2025, 3, 19))
    today = localdate()
    delivery_invoices = _delivery_invoices(today)

    pending_deposits = Invoice.objects.filter(payment_date__isnull=False, deposit_date__isnull=True,
                                              payment_date__gte=start_date).select_related('customer')

    future_deposits = pending_deposits.filter(payment_date__gt=today)
    current_pending_deposits = pending_deposits.exclude(id__in=future_deposits.values('id'))

    total_pending_deposit = current_pending_deposits.aggregate(Sum('total_price'))['total_price__sum'] or 0

    payment_type_totals = current_pending_deposits.values('payment_method').annotate(total=Sum('total_price'))
    future_payment_type_totals = future_deposits.values('payment_method').annotate(total=Sum('total_price'))
    payment_totals_dict = {entry['payment_method']: entry['total'] for entry in payment_type_totals}
    future_payment_totals_dict = {entry['payment_method']: entry['total'] for entry in future_payment_type_totals}

    # Calculate counts for each payment method (group by cheque_detail for cheques)
    payment_counts_dict = {}
    future_payment_counts_dict = {}
    
    for payment_method in set(payment_totals_dict.keys()) | set(future_payment_totals_dict.keys()):
        # For current pending deposits
        current_invoices = current_pending_deposits.filter(payment_method=payment_method)
        if payment_method == 'cheque':
            # Count unique cheque details
            payment_counts_dict[payment_method] = len(set(current_invoices.exclude(cheque_detail__isnull=True).values_list('cheque_detail', flat=True)))
        else:
            # Count individual invoices for non-cheque payments
            payment_counts_dict[payment_method] = current_invoices.count()
        
        # For future deposits
        future_invoices = future_deposits.filter(payment_method=payment_method)
        if payment_method == 'cheque':
            # Count unique cheque details
            future_payment_counts_dict[payment_method] = len(set(future_invoices.exclude(cheque_detail__isnull=True).values_list('cheque_detail', flat=True)))
        else:
            # Count individual invoices for non-cheque payments
            future_payment_counts_dict[payment_method] = future_invoices.count()

    modified_invoices = []
    for invoice in delivery_invoices:
        modified_invoices.append({
            'id': invoice.pk,
            'delivery_date': invoice.delivery_date,
            'number': invoice.number,
            'customer': invoice.customer,
            'deliveryman': invoice.deliveryman,
            'salesman': invoice.salesman,
            'total_price': invoice.total_price,
            'items': _delivery_items(invoice),
        })

    modified_invoices.sort(key=lambda invoice: invoice['delivery_date'] is not None)

    return render(request, 'invoice/home.html', {
        'delivery_invoices': modified_invoices,
        'undelivered_count': sum(invoice['delivery_date'] is None for invoice in modified_invoices),
        'delivered_today_count': sum(invoice['delivery_date'] == today for invoice in modified_invoices),
        'pending_deposits': pending_deposits,
        'total_pending_deposit': total_pending_deposit,
        'payment_totals_dict': payment_totals_dict,
        'future_payment_totals_dict': future_payment_totals_dict,
        'payment_counts_dict': payment_counts_dict,
        'future_payment_counts_dict': future_payment_counts_dict,
    })


@staff_member_required
def sales_data(request):
    """API endpoint providing sales analytics data for charts and reports."""
    try:
        if not Invoice.objects.exists():
            return JsonResponse({"error": "No invoices found"}, status=400)

        current_date = now()
        last_month = (current_date.month - 1) or 12
        last_month_year = current_date.year if current_date.month > 1 else current_date.year - 1
        last_month_name = calendar.month_name[last_month]

        sales_per_month = (
            Invoice.objects
                .annotate(month=ExtractMonth('delivery_date'), year=ExtractYear('delivery_date'))
                .exclude(month__isnull=True)
                .exclude(month=1, year=2025)
                .values('month')
                .annotate(total_sales=Sum('total_price'))
                .order_by('month')
        )

        sales_by_salesman = (
            Invoice.objects
                .filter(delivery_date__month=last_month, delivery_date__year=last_month_year)
                .select_related('salesman')
                .values('salesman__code')
                .annotate(total_sales=Sum('total_price'))
                .exclude(salesman__code="Lafarge")
                .exclude(salesman__code="DS/MM/AC")
                .exclude(salesman__code="KK")
                .order_by('-total_sales')
        )

        data = {
            "sales_per_month": list(sales_per_month),
            "sales_by_salesman": list(sales_by_salesman),
            "last_month_name": last_month_name,
        }

        return JsonResponse(data)

    except Exception as e:
        logger.error(f"Error in sales_data: {str(e)}")
        return JsonResponse({"error": str(e)}, status=500)


@staff_member_required
def product_insights_data(request):
    """API endpoint providing product sales analytics for the previous month."""
    try:
        current_date = now()
        last_month = (current_date.month - 1) or 12
        last_month_year = current_date.year if current_date.month > 1 else current_date.year - 1
        last_month_name = calendar.month_name[last_month]

        # Get all invoice items for the specified month
        from ..models import InvoiceItem
        invoice_items = (
            InvoiceItem.objects
                .filter(invoice__delivery_date__month=last_month, invoice__delivery_date__year=last_month_year)
                .select_related('product')
                .values('product__name', 'sum_price')
        )

        # Group by cleaned product name (without lot numbers)
        grouped_products = defaultdict(float)
        for item in invoice_items:
            if item['product__name']:
                clean_name = re.sub(r"\s*\(Lot\s*no\.?:?\s*[A-Za-z0-9-]+\)", "", item['product__name'])
                grouped_products[clean_name] += float(item['sum_price'] or 0)

        # Convert to list format and sort by revenue
        product_sales = [
            {'invoiceitem__product__name': name, 'total_revenue': revenue}
            for name, revenue in sorted(grouped_products.items(), key=lambda x: x[1], reverse=True)
        ]

        data = {
            "product_sales": product_sales,
            "last_month_name": last_month_name,
        }
        return JsonResponse(data)

    except Exception as e:
        logger.error(f"Error in product_insights_data: {str(e)}")
        return JsonResponse({"error": str(e)}, status=500)
