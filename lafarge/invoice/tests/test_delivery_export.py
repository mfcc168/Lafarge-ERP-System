from datetime import timedelta
from decimal import Decimal
from html.parser import HTMLParser
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils.timezone import localdate
from openpyxl import load_workbook

from ..models import Customer, Invoice, InvoiceItem, Product


class DeliveryCheckboxParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.deliveries = {}

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'input' and attrs.get('name') == 'invoice_ids':
            self.deliveries[attrs['value']] = 'checked' in attrs


class DeliveryExportTest(TestCase):
    def setUp(self):
        self.staff = get_user_model().objects.create_user(username='staff', is_staff=True)
        self.client.force_login(self.staff)
        self.url = reverse('export_today_deliveries')
        self.today = localdate()
        self.customer = Customer.objects.create(
            name='Customer One', care_of='Central',
            address='Flat A\n123 Test Road', delivery_address='Different delivery address',
        )
        self.other_customer = Customer.objects.create(name='Customer Two', address='456 Test Road')
        self.invoice = Invoice.objects.create(
            number='10001', customer=self.customer, delivery_date=self.today,
        )
        self.other_invoice = Invoice.objects.create(
            number='10002', customer=self.other_customer, delivery_date=self.today,
        )
        self.yesterday_invoice = Invoice.objects.create(
            number='10003', customer=self.customer, delivery_date=self.today - timedelta(days=1),
        )
        self.undelivered_invoice = Invoice.objects.create(
            number='10005', customer=self.other_customer,
        )
        self.future_invoice = Invoice.objects.create(
            number='10006', customer=self.customer, delivery_date=self.today + timedelta(days=1),
        )
        tablets = self.create_product('Tablets (Lot no. ABC-1)')
        other_lot = self.create_product('Tablets (Lot no. ABC-2)')
        syrup = self.create_product('Syrup')
        self.add_item(self.invoice, tablets, '2.0')
        self.add_item(self.invoice, other_lot, '3.0')
        self.add_item(self.invoice, syrup, '1.0')
        self.add_item(self.other_invoice, syrup, '4.0')
        self.add_item(self.yesterday_invoice, syrup, '5.0')
        self.add_item(self.undelivered_invoice, syrup, '6.0')

    def create_product(self, name):
        return Product.objects.create(name=name, quantity=Decimal('100.0'), price=Decimal('5.00'))

    def add_item(self, invoice, product, quantity):
        InvoiceItem.objects.create(invoice=invoice, product=product, quantity=Decimal(quantity))

    def export_sheet(self, invoice_ids):
        response = self.client.get(self.url, {'invoice_ids': invoice_ids})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn(f'todays-deliveries-{self.today.isoformat()}.xlsx', response['Content-Disposition'])
        workbook = load_workbook(BytesIO(response.content))
        self.addCleanup(workbook.close)
        return workbook.active

    def test_only_undelivered_invoices_start_selected_on_the_home_page(self):
        response = self.client.get(reverse('home'))
        self.assertEqual(response.status_code, 200)
        parser = DeliveryCheckboxParser()
        parser.feed(response.content.decode())
        self.assertEqual(parser.deliveries, {
            str(self.invoice.pk): False,
            str(self.other_invoice.pk): False,
            str(self.undelivered_invoice.pk): True,
        })
        self.assertEqual(list(parser.deliveries), [
            str(self.undelivered_invoice.pk), str(self.invoice.pk), str(self.other_invoice.pk),
        ])
        self.assertContains(response, 'Not delivered')
        self.assertEqual(response.context['undelivered_count'], 1)
        self.assertEqual(response.context['delivered_today_count'], 2)
        selected_ids = [invoice_id for invoice_id, selected in parser.deliveries.items() if selected]
        self.assertEqual(list(self.export_sheet(selected_ids).values), [
            ('Customer', 'Address', 'Items'),
            ('Customer Two', '456 Test Road', 'Syrup 6.0'),
        ])

    def test_undelivered_invoices_stay_together_before_delivered_invoices(self):
        other_pending = Invoice.objects.create(number='10007', customer=self.customer)
        response = self.client.get(reverse('home'))
        parser = DeliveryCheckboxParser()
        parser.feed(response.content.decode())
        self.assertEqual(list(parser.deliveries), [
            str(self.undelivered_invoice.pk), str(other_pending.pk),
            str(self.invoice.pk), str(self.other_invoice.pk),
        ])
        self.assertEqual(parser.deliveries[str(other_pending.pk)], True)
        self.assertEqual(response.context['undelivered_count'], 2)

    def test_delivered_only_table_starts_with_nothing_selected(self):
        self.undelivered_invoice.delete()
        response = self.client.get(reverse('home'))
        parser = DeliveryCheckboxParser()
        parser.feed(response.content.decode())
        self.assertEqual(parser.deliveries, {
            str(self.invoice.pk): False,
            str(self.other_invoice.pk): False,
        })
        self.assertEqual(response.context['undelivered_count'], 0)

    def test_selected_undelivered_invoice_can_be_exported(self):
        sheet = self.export_sheet([self.undelivered_invoice.pk])
        self.assertEqual(list(sheet.values), [
            ('Customer', 'Address', 'Items'),
            ('Customer Two', '456 Test Road', 'Syrup 6.0'),
        ])

    def test_pending_delivery_is_shown_and_exportable_when_nothing_was_delivered_today(self):
        Invoice.objects.filter(delivery_date=self.today).update(delivery_date=self.today - timedelta(days=1))
        response = self.client.get(reverse('home'))
        self.assertEqual(response.status_code, 200)
        parser = DeliveryCheckboxParser()
        parser.feed(response.content.decode())
        self.assertEqual(parser.deliveries, {str(self.undelivered_invoice.pk): True})
        self.assertEqual(response.context['delivered_today_count'], 0)
        self.assertEqual(list(self.export_sheet([self.undelivered_invoice.pk]).values)[1], (
            'Customer Two', '456 Test Road', 'Syrup 6.0',
        ))

    def test_export_has_only_requested_columns_and_one_row_per_displayed_item(self):
        sheet = self.export_sheet([self.invoice.pk, self.other_invoice.pk])
        self.assertEqual(list(sheet.values), [
            ('Customer', 'Address', 'Items'),
            ('Customer One (Central)', 'Flat A\n123 Test Road', 'Tablets 5.0'),
            ('Customer One (Central)', 'Flat A\n123 Test Road', 'Syrup 1.0'),
            ('Customer Two', '456 Test Road', 'Syrup 4.0'),
        ])

    def test_export_sums_lot_quantities_without_changing_the_product_name(self):
        for lot, quantity in (('ABC-1', '30.0'), ('ABC-2', '15.0')):
            self.add_item(
                self.invoice,
                self.create_product(f'Procough Tab. 60mg (Lot no. {lot})'),
                quantity,
            )
        sheet = self.export_sheet([self.invoice.pk])
        self.assertEqual(list(sheet.values)[-1], (
            'Customer One (Central)', 'Flat A\n123 Test Road', 'Procough Tab. 60mg 45.0',
        ))

    def test_export_includes_only_selected_deliveries_even_for_the_same_customer(self):
        Invoice.objects.create(number='10004', customer=self.customer, delivery_date=self.today)
        sheet = self.export_sheet([self.invoice.pk])
        self.assertEqual(sheet.max_row, 3)
        self.assertTrue(all(row[0] == 'Customer One (Central)' for row in list(sheet.values)[1:]))

    def test_duplicate_selection_does_not_repeat_rows_or_include_other_dates(self):
        sheet = self.export_sheet([
            self.other_invoice.pk, self.other_invoice.pk, self.yesterday_invoice.pk, self.future_invoice.pk,
        ])
        self.assertEqual(list(sheet.values), [
            ('Customer', 'Address', 'Items'),
            ('Customer Two', '456 Test Road', 'Syrup 4.0'),
        ])

    def test_empty_invalid_and_outdated_selections_do_not_export_all_deliveries(self):
        for invoice_ids in ([], ['invalid'], ['-1'], ['9' * 30], [self.yesterday_invoice.pk], [self.future_invoice.pk], [999999]):
            with self.subTest(invoice_ids=invoice_ids):
                response = self.client.get(self.url, {'invoice_ids': invoice_ids})
                self.assertEqual(response.status_code, 400)
                self.assertNotIn('Content-Disposition', response)

    def test_customer_address_and_item_text_are_preserved_as_text(self):
        self.customer.name = '=SUM(1,2)'
        self.customer.care_of = None
        self.customer.address = '=1+1\n中文地址'
        self.customer.save()
        self.invoice.invoiceitem_set.all().delete()
        self.add_item(self.invoice, self.create_product('=1+1'), '2.0')
        sheet = self.export_sheet([self.invoice.pk])
        self.assertEqual(list(sheet.values)[1], ('=SUM(1,2)', '=1+1\n中文地址', '=1+1 2.0'))
        self.assertTrue(all(cell.data_type == 's' for cell in sheet[2]))

    def test_selected_delivery_without_items_is_retained(self):
        invoice = Invoice.objects.create(number='10004', customer=self.customer, delivery_date=self.today)
        sheet = self.export_sheet([invoice.pk])
        self.assertEqual(list(sheet.values)[1], ('Customer One (Central)', 'Flat A\n123 Test Road', None))

    def test_export_requires_a_staff_user(self):
        self.client.logout()
        response = self.client.get(self.url, {'invoice_ids': [self.invoice.pk]})
        self.assertEqual(response.status_code, 302)
        user = get_user_model().objects.create_user(username='regular')
        self.client.force_login(user)
        response = self.client.get(self.url, {'invoice_ids': [self.invoice.pk]})
        self.assertEqual(response.status_code, 302)

    def test_export_accepts_only_get_requests(self):
        response = self.client.post(self.url, {'invoice_ids': [self.invoice.pk]})
        self.assertEqual(response.status_code, 405)
